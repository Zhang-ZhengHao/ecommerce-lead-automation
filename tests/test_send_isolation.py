from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

import services.excel as excel
from services.classifier import classify_demo, process_dataframe, result_column_name
from services.excel import export_result_workbook
from services.policy_defaults import DEFAULT_FORBIDDEN_TERMS
from services.workflow import update_final_reply


def _processed() -> pd.DataFrame:
    return process_dataframe(
        pd.DataFrame({"客户留言": ["想询价", "随便看看"]}),
        "客户留言",
        classify_demo,
    )


def _sendable_export(frame: pd.DataFrame) -> bytes:
    """Call the planned public send-only export interface.

    Keeping the assertion here turns a missing interface into a useful RED
    failure instead of a test-collection error.
    """

    exporter = getattr(excel, "export_sendable_workbook", None)
    assert exporter is not None, "缺少可发送导出接口 export_sendable_workbook"
    return exporter(frame)


def _sheet(payload: bytes, name: str) -> tuple[list[object], list[tuple[object, ...]]]:
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=False)
    try:
        sheet = workbook[name]
        rows = list(sheet.iter_rows(values_only=True))
        return list(rows[0]), [tuple(row) for row in rows[1:]]
    finally:
        workbook.close()


def _column(headers: list[object], name: str) -> int:
    return headers.index(name)


def test_processing_leaves_final_reply_empty_until_manual_confirmation():
    result = _processed()

    final_column = result_column_name(result, "最终回复")
    status_column = result_column_name(result, "回复状态")

    assert result.loc[0, final_column] == ""
    assert result.loc[0, status_column] == "待确认"


def test_unconfirmed_rows_are_absent_from_sendable_export():
    result = _processed()

    headers, rows = _sheet(_sendable_export(result), "可发送")

    assert "原始行号" in headers
    assert rows == []


def test_export_rechecks_unsafe_final_and_routes_row_to_review():
    result = _processed()
    final_column = result_column_name(result, "最终回复")
    status_column = result_column_name(result, "回复状态")
    result.loc[0, final_column] = "保证今天发货"
    result.loc[0, status_column] = "已确认"

    all_headers, _ = _sheet(export_result_workbook(result), "全部结果")
    review_headers, review_rows = _sheet(export_result_workbook(result), "待复核")
    _, sendable_rows = _sheet(_sendable_export(result), "可发送")

    all_row_number = _column(all_headers, "原始行号")
    review_row_number = _column(review_headers, "原始行号")
    assert 1 in [row[review_row_number] for row in review_rows]
    assert 1 not in [row[all_row_number] for row in sendable_rows]


def test_builtin_policy_term_is_blocked_by_sendable_export():
    result = _processed()
    final_column = result_column_name(result, "最终回复")
    status_column = result_column_name(result, "回复状态")
    result.loc[0, final_column] = DEFAULT_FORBIDDEN_TERMS[0]
    result.loc[0, status_column] = "已确认"

    _, rows = _sheet(_sendable_export(result), "可发送")

    assert rows == []


def test_safe_final_is_sendable_even_when_ai_draft_is_unsafe():
    result = _processed()
    draft_column = result_column_name(result, "客户回复")
    final_column = result_column_name(result, "最终回复")
    status_column = result_column_name(result, "回复状态")
    result.loc[0, draft_column] = "保证今天发货"
    result.loc[0, final_column] = "我们先核实订单信息，再回复您。"
    result.loc[0, status_column] = "已确认"

    headers, rows = _sheet(_sendable_export(result), "可发送")

    row_number = _column(headers, "原始行号")
    assert 1 in [row[row_number] for row in rows]


def test_editing_only_ai_draft_does_not_make_row_sendable():
    result = _processed()
    draft_column = result_column_name(result, "客户回复")
    result.loc[0, draft_column] = "人工改过的草稿，但还没有确认。"

    _, rows = _sheet(_sendable_export(result), "可发送")

    assert rows == []


def test_all_results_are_watermarked_and_sendable_is_strict_subset():
    result = _processed()
    result = update_final_reply(result, 0, "人工确认：请补充具体规格。")

    all_headers, all_rows = _sheet(export_result_workbook(result), "全部结果")
    send_headers, send_rows = _sheet(_sendable_export(result), "可发送")

    purpose_all = _column(all_headers, "导出用途")
    purpose_send = _column(send_headers, "导出用途")
    row_all = _column(all_headers, "原始行号")
    row_send = _column(send_headers, "原始行号")
    all_ids = {row[row_all] for row in all_rows}
    send_ids = {row[row_send] for row in send_rows}

    assert all(row[purpose_all] == "草稿-仅供复核" for row in all_rows)
    assert all(row[purpose_send] == "可发送-已人工确认" for row in send_rows)
    assert send_ids == {1}
    assert send_ids < all_ids
