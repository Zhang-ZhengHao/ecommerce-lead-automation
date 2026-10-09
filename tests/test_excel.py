import hashlib
from io import BytesIO

from openpyxl import load_workbook
import pandas as pd
import pytest

from services.excel import (
    MAX_FILE_BYTES,
    MAX_ROWS,
    SpreadsheetError,
    build_template_dataframe,
    build_demo_dataframe,
    content_signature,
    detect_message_column,
    export_result_workbook,
    export_template_xlsx,
    export_xlsx,
    read_table,
)


def test_content_signature_uses_sha256_for_session_deduplication():
    payload = b"synthetic upload"

    assert content_signature(payload) == hashlib.sha256(payload).hexdigest()


def test_read_csv_supports_gb18030_and_preserves_columns():
    source = "客户留言,来源\n你们这个多少钱？,抖音\n我想做代理,小红书\n".encode("gb18030")

    result = read_table(source, "leads.csv")

    assert list(result.columns) == ["客户留言", "来源"]
    assert result.iloc[0]["客户留言"] == "你们这个多少钱？"
    assert result.iloc[1]["来源"] == "小红书"


def test_read_xlsx_uses_first_sheet_and_export_round_trips():
    original = pd.DataFrame(
        {"客户留言": ["报价是多少"], "来源": ["网站"]}
    )
    payload = export_xlsx(original)

    result = read_table(payload, "result.xlsx")

    pd.testing.assert_frame_equal(result, original)


def test_read_table_rejects_unsupported_and_oversized_files():
    with pytest.raises(SpreadsheetError, match="xlsx/csv"):
        read_table(b"anything", "leads.txt")

    with pytest.raises(SpreadsheetError, match="10 MB"):
        read_table(b"0" * (MAX_FILE_BYTES + 1), "leads.csv")


def test_read_table_rejects_overlarge_row_count_before_processing():
    payload = ("留言\n" + "正常\n" * (MAX_ROWS + 1)).encode("utf-8")

    with pytest.raises(SpreadsheetError, match="5,000"):
        read_table(payload, "leads.csv")


def test_read_table_makes_normalised_duplicate_headers_unique():
    payload = " 客户留言 ,客户留言,客户留言\n第一条,第二条,第三条\n".encode("utf-8")

    result = read_table(payload, "duplicate.csv")

    assert list(result.columns) == ["客户留言", "客户留言（2）", "客户留言（3）"]


def test_export_xlsx_escapes_formula_like_strings():
    payload = export_xlsx(
        pd.DataFrame(
            {
                "留言": ['=HYPERLINK("https://evil.example","点击")', "正常"],
                "回复": ["+cmd", "\t@mention"],
                '=HYPERLINK("https://evil.example","表头")': ["safe", "safe"],
            }
        )
    )

    workbook = pd.read_excel(BytesIO(payload), engine="openpyxl")

    assert workbook.iloc[0, 0].startswith("'=")
    assert workbook.iloc[0, 1].startswith("'+")
    assert workbook.iloc[1, 1].startswith("'\t@")
    assert workbook.columns[2].startswith("'=")


def test_demo_dataframe_is_ready_for_the_full_flow():
    result = build_demo_dataframe()

    assert list(result.columns) == ["客户留言", "联系方式", "来源"]
    assert len(result) >= 3
    assert result["客户留言"].notna().all()


def test_demo_dataframe_includes_personalisation_examples():
    result = build_demo_dataframe()
    messages = result["客户留言"].tolist()

    assert any("100件" in message and "预算" in message for message in messages)
    assert any("订单号" in message and "坏了" in message for message in messages)


def test_template_dataframe_has_standard_input_columns_and_synthetic_example():
    result = build_template_dataframe()

    assert list(result.columns) == ["客户留言", "联系方式", "来源", "时间"]
    assert len(result) == 1
    assert "示例" in result.iloc[0]["客户留言"]


def test_template_workbook_contains_input_and_instruction_sheets():
    payload = export_template_xlsx()

    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=False)
    try:
        assert workbook.sheetnames == ["填写模板", "使用说明"]
        headers = next(workbook["填写模板"].iter_rows(values_only=True))
        assert list(headers) == ["客户留言", "联系方式", "来源", "时间"]
        instruction_values = [
            cell.value
            for row in workbook["使用说明"].iter_rows()
            for cell in row
            if cell.value is not None
        ]
        assert any("客户留言" in str(value) for value in instruction_values)
        assert any("上传" in str(value) for value in instruction_values)
        assert any("删除" in str(value) or "替换" in str(value) for value in instruction_values)
    finally:
        workbook.close()


def test_detect_message_column_prefers_high_priority_aliases():
    frame = pd.DataFrame(columns=["来源", "咨询内容", "客户留言"])

    assert detect_message_column(frame) == "客户留言"


@pytest.mark.parametrize(
    ("columns", "expected"),
    [
        (["来源", "留言内容"], "留言内容"),
        (["来源", "Customer_Message"], "Customer_Message"),
        (["来源", " customer-message "], " customer-message "),
        (["来源", "INQUIRY"], "INQUIRY"),
    ],
)
def test_detect_message_column_accepts_common_aliases_and_columns_iterable(
    columns, expected
):
    assert detect_message_column(columns) == expected


def test_detect_message_column_does_not_treat_metadata_as_message():
    frame = pd.DataFrame(columns=["留言时间", "来源", "联系人"])

    assert detect_message_column(frame) is None


def test_export_result_workbook_contains_four_filtered_sheets():
    frame = pd.DataFrame(
        {
            "客户留言": ["想询价", "随便看看", "订单故障"],
            "意向等级": ["高", "低", "中"],
            "客户回复": ["请问规格？", "欢迎了解", "请提供订单号"],
            "处理状态": ["成功", "成功", "模型失败"],
        }
    )

    payload = export_result_workbook(frame)

    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=False)
    try:
        assert workbook.sheetnames == ["全部结果", "高意向", "待复核", "使用说明"]
        assert workbook["全部结果"].max_row == 4
        assert workbook["高意向"].max_row == 2
        assert workbook["待复核"].max_row == 2
        assert list(
            next(workbook["全部结果"].iter_rows(values_only=True))
        ) == ["原始行号", "导出用途", *frame.columns]
        instruction_values = [
            cell.value
            for row in workbook["使用说明"].iter_rows()
            for cell in row
            if cell.value is not None
        ]
        assert any("高意向" in str(value) for value in instruction_values)
    finally:
        workbook.close()


def test_export_result_workbook_keeps_headers_when_filters_are_empty():
    frame = pd.DataFrame(
        {
            "客户留言": ["暂不考虑"],
            "意向等级": ["低"],
            "处理状态": ["成功"],
        }
    )

    payload = export_result_workbook(frame)

    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=False)
    try:
        for sheet_name in ("高意向", "待复核"):
            sheet = workbook[sheet_name]
            assert sheet.max_row == 1
            assert list(next(sheet.iter_rows(values_only=True))) == [
                "原始行号",
                "导出用途",
                *frame.columns,
            ]
    finally:
        workbook.close()


def test_export_result_workbook_escapes_formula_like_values():
    frame = pd.DataFrame(
        {
            "客户留言": ['=HYPERLINK("https://evil.example","点击")'],
            "意向等级": ["高"],
            "客户回复": ["+cmd"],
            "处理状态": ["成功"],
        }
    )

    payload = export_result_workbook(frame)

    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=False)
    try:
        for sheet_name in ("全部结果", "高意向"):
            values = list(workbook[sheet_name].iter_rows(values_only=True))
            headers = list(values[0])
            row = values[1]
            assert row[headers.index("客户留言")].startswith("'=")
            assert row[headers.index("客户回复")].startswith("'+")
    finally:
        workbook.close()


def test_export_result_workbook_uses_collision_safe_result_columns():
    frame = pd.DataFrame(
        {
            "客户留言": ["高意向", "失败"],
            "意向等级": ["原始值1", "原始值2"],
            "意向等级（AI）": ["高", "低"],
            "处理状态": ["原始状态1", "原始状态2"],
            "处理状态（AI）": ["成功", "模型失败"],
        }
    )
    frame.attrs["commerce_lead_result_columns"] = {
        "意向等级": "意向等级（AI）",
        "处理状态": "处理状态（AI）",
    }

    payload = export_result_workbook(frame)

    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=False)
    try:
        assert workbook["高意向"].max_row == 2
        assert workbook["待复核"].max_row == 2
    finally:
        workbook.close()


def test_export_result_workbook_uses_last_collision_column_when_attrs_are_missing():
    frame = pd.DataFrame(
        {
            "客户留言": ["第一条", "第二条"],
            "意向等级": ["原始1", "原始2"],
            "意向等级（AI）": ["低", "低"],
            "意向等级（AI2）": ["高", "低"],
            "处理状态": ["原始1", "原始2"],
            "处理状态（AI）": ["成功", "成功"],
            "处理状态（AI2）": ["成功", "模型失败"],
        }
    )

    payload = export_result_workbook(frame)

    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=False)
    try:
        assert workbook["高意向"].max_row == 2
        assert workbook["待复核"].max_row == 2
    finally:
        workbook.close()


def test_export_result_workbook_contains_draft_final_and_status_columns():
    frame = pd.DataFrame(
        {
            "客户留言": ["询价"],
            "意向等级": ["高"],
            "处理状态": ["成功"],
            "客户回复": ["AI 草稿"],
            "最终回复": ["人工确认"],
            "回复状态": ["已确认"],
        }
    )
    workbook = load_workbook(
        BytesIO(export_result_workbook(frame)), read_only=True, data_only=False
    )
    try:
        expected_headers = ["原始行号", "导出用途", *frame.columns]
        assert list(next(workbook["全部结果"].iter_rows(values_only=True))) == expected_headers
        assert list(next(workbook["高意向"].iter_rows(values_only=True))) == expected_headers
    finally:
        workbook.close()


def test_export_escapes_formula_like_final_reply():
    frame = pd.DataFrame(
        {
            "客户留言": ["询价"],
            "意向等级": ["高"],
            "处理状态": ["成功"],
            "客户回复": ["草稿"],
            "最终回复": ['=HYPERLINK("https://evil.example","点我")'],
            "回复状态": ["已确认"],
        }
    )
    workbook = load_workbook(
        BytesIO(export_result_workbook(frame)), read_only=True, data_only=False
    )
    try:
        headers = list(next(workbook["全部结果"].iter_rows(values_only=True)))
        final_column = headers.index("最终回复") + 1
        assert workbook["全部结果"].cell(2, final_column).value.startswith("'=")
    finally:
        workbook.close()


def test_export_result_workbook_review_includes_reply_status_without_processing_failure():
    frame = pd.DataFrame(
        {
            "客户留言": ["已确认", "待复核"],
            "意向等级": ["高", "高"],
            "处理状态": ["成功", "成功"],
            "客户回复": ["草稿1", "草稿2"],
            "最终回复": ["最终1", "最终2"],
            "回复状态": ["已确认", "待复核"],
        }
    )
    workbook = load_workbook(
        BytesIO(export_result_workbook(frame)), read_only=True, data_only=False
    )
    try:
        assert workbook["待复核"].max_row == 2
        assert workbook["高意向"].max_row == 3
    finally:
        workbook.close()


def test_export_result_workbook_instructions_describe_reply_fields():
    frame = pd.DataFrame({"客户留言": ["询价"], "处理状态": ["成功"]})
    workbook = load_workbook(
        BytesIO(export_result_workbook(frame)), read_only=True, data_only=False
    )
    try:
        values = [
            str(cell.value)
            for row in workbook["使用说明"].iter_rows()
            for cell in row
            if cell.value is not None
        ]
        text = " ".join(values)
        assert "客户回复" in text and "AI 草稿" in text
        assert "最终回复" in text and "人工确认" in text
        assert all(status in text for status in ("待确认", "已确认", "待复核"))
    finally:
        workbook.close()


def test_watermark_export_preserves_original_row_number_for_filtered_frame():
    frame = pd.DataFrame({"客户留言": ["第一条", "第二条", "第三条"]})
    payload = export_xlsx(frame.iloc[[2]], purpose="草稿-仅供复核")
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=False)
    try:
        rows = list(workbook["线索处理结果"].iter_rows(values_only=True))
        headers, row = rows[0], rows[1]
        assert row[headers.index("原始行号")] == 3
    finally:
        workbook.close()
