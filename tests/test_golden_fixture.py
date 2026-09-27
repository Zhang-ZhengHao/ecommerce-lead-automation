import json
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from services.classifier import classify_demo, process_dataframe
from services.excel import detect_message_column, export_result_workbook, read_table


FIXTURE = Path(__file__).parent / "fixtures" / "golden.xlsx"
META = FIXTURE.with_suffix(".meta.json")
EXPECTED_IDS = [f"T{number:02d}" for number in range(1, 16)]


def _golden_result():
    source = read_table(FIXTURE.read_bytes(), "golden.xlsx")
    message_column = detect_message_column(source)
    assert message_column == "咨询内容"
    return source, process_dataframe(source, message_column, classify_demo)


def _golden_meta():
    return json.loads(META.read_text(encoding="utf-8"))


def test_golden_metadata_declares_expectation_provenance_and_coverage():
    meta = _golden_meta()
    fixture_ids = set(EXPECTED_IDS)
    verified_ids = set(meta["verified_ids"])
    unverified_ids = set(meta["unverified_ids"])

    assert meta["fixture"] == FIXTURE.name
    assert meta["expectation_source"] == "author_prior"
    assert meta["last_reviewed"] == "2026-09-21"
    assert verified_ids
    assert verified_ids.isdisjoint(unverified_ids)
    assert verified_ids | unverified_ids == fixture_ids
    source_by_id = meta["expectation_source_by_id"]
    assert set(source_by_id) == fixture_ids
    assert {
        test_id for test_id, source in source_by_id.items() if source == "unverified"
    } == unverified_ids
    assert all(source_by_id[test_id] == "author_prior" for test_id in verified_ids)
    assert set(meta["expected"]["intent"]) == fixture_ids
    assert set(meta["expected"]["category"]) == fixture_ids
    assert set(meta["expected"]["processing_status"]) == fixture_ids
    assert set(meta["expected"]["reply_status"]) == fixture_ids
    assert set(meta["unverified_reasons"]) == unverified_ids


def test_golden_fixture_reads_first_sheet_and_preserves_test_ids():
    source, result = _golden_result()

    assert len(source) == 15
    assert source["测试编号"].tolist() == EXPECTED_IDS
    assert result["测试编号"].tolist() == EXPECTED_IDS
    assert list(result["联系方式"]) == list(source["联系方式"])
    assert list(result["来源"]) == list(source["来源"])


def test_golden_fixture_matches_demo_classification_baseline():
    _, result = _golden_result()
    meta = _golden_meta()
    expected_intent = meta["expected"]["intent"]
    expected_category = meta["expected"]["category"]
    by_id = result.set_index("测试编号")

    for test_id in meta["verified_ids"]:
        assert by_id.at[test_id, "意向等级"] == expected_intent[test_id]
        assert by_id.at[test_id, "客户类型"] == expected_category[test_id]

    # Unverified rows intentionally do not participate in classification
    # accuracy assertions; their current outputs are retained in metadata for
    # audit only until a customer provides labeled examples.
    unverified = by_id.loc[meta["unverified_ids"]]
    assert len(unverified) == len(meta["unverified_ids"])
    assert {
        "测试编号",
        "意向等级",
        "客户类型",
        "处理状态",
        "回复状态",
    }.issubset(result.columns)
    assert unverified["意向等级"].map(lambda value: isinstance(value, str)).all()
    assert unverified["客户类型"].map(lambda value: isinstance(value, str)).all()

    # Processing/reply statuses are exact regression expectations only for
    # verified rows. Unverified rows retain only a legal-status/format check.
    verified = by_id.loc[meta["verified_ids"]]
    expected_processing = meta["expected"]["processing_status"]
    expected_reply = meta["expected"]["reply_status"]
    assert verified["处理状态"].to_dict() == {
        test_id: expected_processing[test_id] for test_id in meta["verified_ids"]
    }
    assert verified["回复状态"].to_dict() == {
        test_id: expected_reply[test_id] for test_id in meta["verified_ids"]
    }
    allowed_processing = {"成功", "空内容", "模型失败"}
    allowed_reply = {"待确认", "已确认", "待复核"}
    assert unverified["处理状态"].map(
        lambda value: isinstance(value, str) and value in allowed_processing
    ).all()
    assert unverified["回复状态"].map(
        lambda value: isinstance(value, str) and value in allowed_reply
    ).all()


def test_golden_fixture_keeps_final_reply_send_track_empty_by_default():
    _, result = _golden_result()
    meta = _golden_meta()
    by_id = result.set_index("测试编号")

    assert result["最终回复"].fillna("").eq("").all()
    expected_reply = meta["expected"]["reply_status"]
    verified = by_id.loc[meta["verified_ids"]]
    assert verified["回复状态"].to_dict() == {
        test_id: expected_reply[test_id] for test_id in meta["verified_ids"]
    }
    unverified = by_id.loc[meta["unverified_ids"]]
    assert unverified["回复状态"].map(
        lambda value: isinstance(value, str)
        and value in {"待确认", "已确认", "待复核"}
    ).all()


def test_golden_fixture_formula_like_message_is_escaped_on_export():
    _, result = _golden_result()
    workbook = load_workbook(
        BytesIO(export_result_workbook(result)), read_only=True, data_only=False
    )
    try:
        sheet = workbook["全部结果"]
        rows = list(sheet.iter_rows(values_only=True))
        headers = list(rows[0])
        id_column = headers.index("测试编号")
        message_column = headers.index("咨询内容")
        row = next(row for row in rows[1:] if row[id_column] == "T12")
        assert str(row[message_column]).startswith("'=")
    finally:
        workbook.close()
