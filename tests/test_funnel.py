import pandas as pd
import pytest
from io import BytesIO
from openpyxl import load_workbook

from services.classifier import classify_demo, process_dataframe, result_column_name
from services.funnel import (
    DEFAULT_FUNNEL_STAGE,
    FUNNEL_STAGES,
    FUNNEL_COLUMN,
    is_valid_funnel_stage,
    normalize_funnel_stage,
)
from services.industry_presets import (
    INDUSTRY_PRESET_OPTIONS,
    get_industry_preset,
)
from services.results import filter_and_sort_results
from services.excel import export_result_workbook
from services.excel import build_demo_dataframe
from services.workflow import (
    ReplyEditError,
    preserve_confirmed_replies,
    update_funnel_stage,
)


def test_funnel_stage_constants_and_normalisation_are_safe_for_imported_values():
    assert FUNNEL_STAGES == ("新线索", "已联系", "待补信息", "已转交")
    assert DEFAULT_FUNNEL_STAGE == "新线索"
    assert is_valid_funnel_stage("已联系")
    assert not is_valid_funnel_stage("已发送")
    assert normalize_funnel_stage("已联系") == "已联系"
    assert normalize_funnel_stage("未知阶段") == DEFAULT_FUNNEL_STAGE
    assert normalize_funnel_stage(None) == DEFAULT_FUNNEL_STAGE


def test_processing_adds_session_funnel_stage_and_routes_invalid_values_to_review():
    source = pd.DataFrame(
        {
            "客户留言": ["想询价", "我想做代理"],
            FUNNEL_COLUMN: ["已联系", "不是合法阶段"],
        }
    )

    result = process_dataframe(source, "客户留言", classify_demo)
    stage_column = result_column_name(result, FUNNEL_COLUMN)
    reply_status_column = result_column_name(result, "回复状态")
    reason_column = result_column_name(result, "判断依据")

    assert result[stage_column].tolist() == ["已联系", DEFAULT_FUNNEL_STAGE]
    assert result[reply_status_column].tolist() == ["待确认", "待复核"]
    assert "线索阶段" in str(result.loc[1, reason_column])


def test_processing_defaults_funnel_stage_when_input_has_no_stage_column():
    source = pd.DataFrame({"客户留言": ["随便看看"]})

    result = process_dataframe(source, "客户留言", classify_demo)

    assert result_column_name(result, FUNNEL_COLUMN) in result.columns
    assert result.iloc[0][result_column_name(result, FUNNEL_COLUMN)] == DEFAULT_FUNNEL_STAGE


def test_processing_routes_invalid_classifier_stage_to_review():
    source = pd.DataFrame({"客户留言": ["询价"]})

    def classifier(_text: str):
        return {
            "意向等级": "高",
            "客户类型": "询价",
            "线索阶段": "模型乱填",
            "客户回复": "请补充规格。",
            "跟进动作": "确认规格",
            "判断依据": "测试",
        }

    result = process_dataframe(source, "客户留言", classifier)
    stage_column = result_column_name(result, FUNNEL_COLUMN)
    status_column = result_column_name(result, "回复状态")

    assert result.loc[0, stage_column] == DEFAULT_FUNNEL_STAGE
    assert result.loc[0, status_column] == "待复核"


def test_invalid_imported_stage_is_explained_even_for_empty_messages():
    result = process_dataframe(
        pd.DataFrame({"客户留言": [""], FUNNEL_COLUMN: ["阶段写错"]}),
        "客户留言",
        classify_demo,
    )
    reason_column = result_column_name(result, "判断依据")

    assert result.loc[0, result_column_name(result, "回复状态")] == "待复核"
    assert "线索阶段无效" in str(result.loc[0, reason_column])


def test_update_funnel_stage_is_session_only_and_validates_stage():
    result = process_dataframe(
        pd.DataFrame({"客户留言": ["想询价"]}), "客户留言", classify_demo
    )

    updated = update_funnel_stage(result, 0, "已转交")

    assert updated is not result
    assert updated.loc[0, result_column_name(updated, FUNNEL_COLUMN)] == "已转交"
    assert result.loc[0, result_column_name(result, FUNNEL_COLUMN)] == DEFAULT_FUNNEL_STAGE
    with pytest.raises(ReplyEditError, match="线索阶段"):
        update_funnel_stage(result, 0, "未知阶段")


def test_reprocessing_same_source_preserves_session_funnel_stage():
    source = pd.DataFrame({"客户留言": ["想询价"]})
    first = process_dataframe(source, "客户留言", classify_demo)
    edited = update_funnel_stage(first, 0, "已转交")
    fresh = process_dataframe(source, "客户留言", classify_demo)

    preserved = preserve_confirmed_replies(
        fresh,
        edited,
        text_column="客户留言",
        previous_text_column="客户留言",
    )

    assert preserved.loc[0, result_column_name(preserved, FUNNEL_COLUMN)] == "已转交"


def test_results_can_filter_by_funnel_stage_without_mutating_source():
    result = process_dataframe(
        pd.DataFrame({"客户留言": ["想询价", "随便看看"]}),
        "客户留言",
        classify_demo,
    )
    stage_column = result_column_name(result, FUNNEL_COLUMN)
    result.loc[1, stage_column] = "已联系"

    filtered = filter_and_sort_results(result, funnel_stage="已联系")

    assert filtered.index.tolist() == [1]
    assert result.index.tolist() == [0, 1]


def test_industry_presets_return_valid_profiles_and_reject_unknown_names():
    assert INDUSTRY_PRESET_OPTIONS == ("电商咨询", "售后", "B2B 询价")
    for name in INDUSTRY_PRESET_OPTIONS:
        profile = get_industry_preset(name)
        assert profile.industry_context
        assert profile.tone in {"专业", "亲切", "简洁", "稳重"}
    with pytest.raises(ValueError, match="行业预设"):
        get_industry_preset("不存在")


def test_result_export_instructions_explain_session_funnel_stages():
    frame = process_dataframe(
        pd.DataFrame({"客户留言": ["想询价"]}), "客户留言", classify_demo
    )

    workbook = load_workbook(
        BytesIO(export_result_workbook(frame)), read_only=True, data_only=False
    )
    try:
        text = " ".join(
            str(cell.value)
            for row in workbook["使用说明"].iter_rows()
            for cell in row
            if cell.value is not None
        )
    finally:
        workbook.close()
    assert "线索阶段" in text
    assert all(stage in text for stage in FUNNEL_STAGES)


def test_demo_contacts_are_explicit_synthetic_ids_not_phone_numbers():
    contacts = build_demo_dataframe()["联系方式"].astype(str).tolist()

    assert contacts == [
        "demo-contact-001",
        "demo-contact-002",
        "demo-contact-003",
        "demo-contact-004",
    ]
    assert not any(contact.isdigit() and len(contact) >= 7 for contact in contacts)
